from .base_loader import BaseLoader
from typing import List, Dict
import logging

logger = logging.getLogger(__name__)


class PPTLoader(BaseLoader):
    def load(self) -> List[Dict]:
        try:
            from pptx import Presentation
        except Exception:
            logger.exception('python-pptx not installed')
            return []

        try:
            prs = Presentation(self.path)
            texts = []
            for i, slide in enumerate(prs.slides, start=1):
                parts = []
                for shape in slide.shapes:
                    if hasattr(shape, 'text'):
                        parts.append(shape.text)
                texts.append({'page': i, 'text': '\n'.join(parts)})
            return texts
        except Exception:
            logger.exception('Failed to read ppt')
            return []
