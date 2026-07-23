from .base_loader import BaseLoader
from typing import List, Dict
import logging

logger = logging.getLogger(__name__)


class ExcelLoader(BaseLoader):
    def load(self) -> List[Dict]:
        # Try openpyxl first, then pandas as fallback
        try:
            from openpyxl import load_workbook
        except Exception:
            try:
                import pandas as pd
            except Exception:
                logger.exception('openpyxl or pandas not installed for Excel loading')
                return []
            else:
                try:
                    df = pd.read_excel(self.path, sheet_name=None)
                    texts = []
                    for sheet, frame in df.items():
                        text = '\n'.join(['\t'.join([str(x) for x in row if pd.notna(x)]) for row in frame.values.tolist()])
                        texts.append({'page': 1, 'text': text})
                    return texts
                except Exception:
                    logger.exception('Failed to read excel with pandas')
                    return []

        try:
            wb = load_workbook(filename=self.path, read_only=True, data_only=True)
            texts = []
            for sheet in wb.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    cols = [str(c) for c in row if c is not None]
                    if cols:
                        rows.append('\t'.join(cols))
                texts.append({'page': 1, 'text': '\n'.join(rows)})
            return texts
        except Exception:
            logger.exception('Failed to read excel with openpyxl')
            return []
